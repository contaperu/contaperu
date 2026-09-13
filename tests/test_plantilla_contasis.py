"""Las columnas del driver CONTASIS contra la plantilla oficial, y sus celdas contra un registro que CONTASIS importó.

Los dos registros son de un contribuyente real y viven fuera de Git, en `tests/fixtures/privado/contasis/`, junto a
las plantillas: sin ellos, esta prueba se salta. De ellos solo se miran formas —cuántas columnas, qué largo, qué
tipo de celda, cómo va una columna vacía—, nunca valores, y los mensajes de error tampoco los muestran.
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path

import pytest

from contaperu.drivers import contasis
from test_snapshot_contasis import CASOS, fila_de

PRIVADO = Path(__file__).parent / "fixtures" / "privado" / "contasis"
PLANTILLAS = {"compra": "FORMATO REGISTRO DE COMPRAS.xlsx", "venta": "FORMATO REGISTRO DE VENTAS.xlsx"}
VALIDADOS = {"compra": "REGISTRO DE COMPRAS - EJEMPLO.xlsx", "venta": "REGISTRO DE VENTAS - EJEMPLO.xlsx"}

pytestmark = pytest.mark.skipif(
    not all((PRIVADO / n).exists() for n in (*PLANTILLAS.values(), *VALIDADOS.values())),
    reason="sin la plantilla y el registro validado de CONTASIS en tests/fixtures/privado/contasis (fuera de Git)")

D = contasis.datos


def _clase_y_largo(formato: str) -> tuple[str, int]:
    """La fila 13 de la plantilla: `dd/mm/yyyy`, `(15,2) NUMERICO`, `02 CARACTERES`, `(10) CARACTER`…"""
    f = " ".join(formato.split())
    if f == "dd/mm/yyyy":
        return D.FECHA, 0
    if f == "(10,4) NUMERICO":
        return D.CAMBIO, 0
    if f == "(5,2) NUMERICO":
        return D.PORCENTAJE, 0
    if re.fullmatch(r"\(\d+,2\) NUMERICO", f):
        return D.IMPORTE, 0
    if f == "1 NUMERICO":
        return D.NUMERO, 1
    m = re.fullmatch(r"\(?(\d+)\)? CARACTER(ES)?", f)
    assert m, f"formato desconocido en la fila 13 de la plantilla: {f!r}"
    return D.TEXTO, int(m.group(1))


@pytest.mark.parametrize("tipo", ["compra", "venta"])
def test_las_columnas_son_las_de_la_plantilla(tipo):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.utils import get_column_letter

    ws = openpyxl.load_workbook(PRIVADO / PLANTILLAS[tipo]).active
    assert ws.title == D.HOJAS[tipo]
    de_la_plantilla = [(get_column_letter(c), *_clase_y_largo(str(ws.cell(13, c).value or "")))
                       for c in range(1, ws.max_column + 1)]
    assert [(letra, clase, largo) for letra, _, clase, largo in D.COLUMNAS[tipo]] == de_la_plantilla


def _forma(valor) -> str | None:
    """None para lo vacío (una celda vacía o solo espacios); si no, la clase de la celda."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return "fecha"
    return "numero" if isinstance(valor, (int, float)) else "texto"


@pytest.mark.parametrize("tipo", ["compra", "venta"])
def test_cada_celda_tiene_la_forma_del_registro_validado(tipo):
    openpyxl = pytest.importorskip("openpyxl")

    ws = openpyxl.load_workbook(PRIVADO / VALIDADOS[tipo]).active
    validadas = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                 for r in range(14, ws.max_row + 1) if ws.cell(r, 1).value is not None]
    nuestras = [fila_de(caso) for caso in CASOS if caso[2] == (tipo == "venta")]
    assert validadas and nuestras
    for i, (letra, _, clase, largo) in enumerate(D.COLUMNAS[tipo]):
        suyas = [f[i] for f in validadas]
        mias = [f.get(letra) for f in nuestras]
        formas = Counter(_forma(v) for v in suyas if _forma(v))
        if formas:
            forma = formas.most_common(1)[0][0]
            assert {_forma(v) for v in mias if _forma(v)} <= {forma}, f"{letra}: el registro validado la escribe como {forma}"
        if clase != D.TEXTO:
            continue
        largos = Counter(len(v) for v in suyas if isinstance(v, str))
        if largos and largos.most_common(1)[0][0] == largo:
            assert all(len(v) == largo for v in mias if isinstance(v, str)), f"{letra}: el registro validado la rellena a {largo}"
        # Espacios de verdad frente a una celda vacía: el registro validado deja algunas con un texto vacío (""), que
        # para CONTASIS es lo mismo que no escribirla.
        if any(isinstance(v, str) and v and not v.strip() for v in suyas):
            assert letra not in D.VACIAS_SIN_ESPACIOS[tipo], f"{letra}: vacía, el registro validado la llena de espacios"
        elif all(v in (None, "") for v in suyas):
            assert letra in D.VACIAS_SIN_ESPACIOS[tipo], f"{letra}: vacía, el registro validado la deja sin nada"
