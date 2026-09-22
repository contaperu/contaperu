"""Excel de asientos para CONCAR: la mecánica del formato, caso por caso.

PEN con IGV → 3 filas; USD → `US`, P, 421202 y T.C. en G con conversión C; boleta → 13/BV sin
IGV; recibo por honorarios → 15 a 424101; nota de crédito invierte; detracción → sub-diario 10
+ AI–AL. Incluye los tres errores que un generador anterior cometía y que aquí se prueban a
propósito: el MM del periodo, los correlativos obligatorios y la cuenta obligatoria."""
import io
import uuid
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from contaperu.pipeline import salida as gen
from contaperu import api
from contaperu.pipeline import preparacion as prep
from contaperu.modelo import Comprobante, Libro
from contaperu import asiento as concar
from contaperu.drivers import concar as driver_concar
from util import comprobante, con_imputaciones, construir_concar, en_secciones

COMPRAS = Libro(ruc="20601111111", razon_social="EMPRESA DE PRUEBA SAC", periodo="202608", tipo="compra")
VENTAS = Libro(ruc="20601111111", razon_social="EMPRESA DE PRUEBA", periodo="202608", tipo="venta")
MES = (date(2026, 8, 1), date(2026, 8, 31))   # limites del periodo 202608 (la fecha va por comprobante)
EMISION, VENCE = date(2026, 8, 11), date(2026, 8, 18)
def configuracion(config_contable: dict | None = None) -> dict:
    """La configuración aplicada para CONCAR (la del entorno sobre la de por defecto) con las imputaciones de las
    pruebas. Se escribe plana y se guarda en su sección (`util.en_secciones`)."""
    return con_imputaciones(prep.config_aplicada(en_secciones(config_contable, "concar"), "concar"))


CONTAB = configuracion(None)
# CONCAR trae su cuenta de gasto (`datos.CUENTAS_POR_DEFECTO`, 2.5), así que para probar que un comprobante SIN
# cuenta se planta hay que decir que no hay ninguna — que es lo que hace quien prefiere que le avisen.
SIN_GASTO = configuracion({"cuentas": {"gasto": ""}})
# El centro también en la X de la línea del gasto, como referencia, además del doble anexo del tercero.
REFERENCIA_EN_X = {"columnas": {"centro_costo": ["centro_costo", "anexo_auxiliar", "anexo_auxiliar_del_tercero"]}}


def cp(**k):
    base = dict(tipo_cp="01", serie="F001", numero="00000123", fecha_emision="2026-08-11", fecha_vencimiento="2026-08-18",
                contraparte_tipo_doc="6", contraparte_doc="20607777773", contraparte_nombre="ELECTROMECANICA DE PRUEBA S.R.L.",
                moneda="PEN", base_gravada="100", igv="18", total="118", concepto="Grillete tipo lira 5/8 para torre de alta tension linea 2",
                cuenta_contable="631101", centro_costo="OBRA01")
    base.update(k)
    return comprobante(**base)


def debe_haber(filas):
    return (sum(Decimal(str(f["O"])) for f in filas if f["N"] == "D"), sum(Decimal(str(f["O"])) for f in filas if f["N"] == "H"))


# ── El asiento ────────────────────────────────────────────────────────────────

def test_factura_pen_con_igv_tres_filas():
    filas = driver_concar.filas_de_comprobante(cp(), CONTAB, MES, "080020")
    assert len(filas) == 3
    gasto, igv, cxp = filas
    assert [f["N"] for f in filas] == ["D", "D", "H"]
    assert [f["K"] for f in filas] == ["631101", "401111", "421201"]
    assert [f["O"] for f in filas] == [100.0, 18.0, 118.0]
    assert [f["Q"] for f in filas] == [100.0, 18.0, 118.0] and all(f["P"] == "" for f in filas)   # soles → Q, P vacío
    assert gasto["M"] == "OBRA01" and igv["M"] == "" and cxp["M"] == ""                           # centro de costo (M) solo en el gasto
    assert cxp["X"] == "OBRA01" and gasto["X"] == "" and igv["X"] == ""                           # … y en el anexo auxiliar del proveedor (doble anexo)
    assert cxp["L"] == "20607777773" and gasto["L"] == "" and igv["L"] == ""                      # anexo (RUC) solo en el proveedor
    for f in filas:
        assert (f["B"], f["C"], f["D"], f["E"], f["G"], f["H"], f["I"], f["J"]) == ("11", "080020", EMISION, "MN", "", "V", "S", EMISION)
        assert (f["R"], f["S"], f["T"], f["U"], f["AO"], f["A"], f["AI"]) == ("FT", "F001-123", EMISION, VENCE, 18, "", "")
        assert len(f) == 41 and list(f.keys()) == driver_concar.datos.COLUMNAS
    # Glosas en MAYÚSCULAS. La principal (F) se corta a 40 y la de detalle (W) a 30, y desde la 2.2 la línea del
    # IGV lleva LA MISMA glosa que las demás: el prefijo `IGV - ` gastaba 6 de esos 30 caracteres.
    assert gasto["F"] == igv["F"] == cxp["F"] == "GRILLETE TIPO LIRA 5/8 PARA TORRE DE ALTA TENSION LINEA 2"[:40].upper()
    assert gasto["W"] == igv["W"] == cxp["W"] == "GRILLETE TIPO LIRA 5/8 PARA TO" and len(gasto["W"]) == 30
    assert debe_haber(filas) == (Decimal("118"), Decimal("118"))


def test_en_soles_el_tipo_de_cambio_no_entra_aunque_el_comprobante_lo_traiga():
    """CONCAR no convierte nada en soles, así que su columna G se queda vacía y la conversión, en la de siempre.

    Importa desde el 22-sep-2026: hay aplicaciones que anotan el T.C. del día también en PEN —STARSOFT lo escribe
    en todas sus filas—, y desde la 2.2 la LÍNEA puede traerlo en soles. Si aquí se mirara solo «viene lleno» en
    vez de la moneda, el Excel de un mes entero cambiaría de columna G y de tipo de conversión sin que nadie lo
    hubiera pedido."""
    con_tc = driver_concar.filas_de_comprobante(cp(tipo_cambio="3.383"), CONTAB, MES, "080001")
    sin_tc = driver_concar.filas_de_comprobante(cp(), CONTAB, MES, "080001")
    assert con_tc[0]["G"] == "" and con_tc[0]["H"] == sin_tc[0]["H"]
    assert [f["G"] for f in con_tc] == [f["G"] for f in sin_tc] == ["", "", ""]
    # Y la fila entera es la misma: el T.C. en soles no mueve ni una celda.
    assert con_tc == sin_tc


def test_dolares_con_tipo_de_cambio_del_comprobante():
    usd = driver_concar.filas_de_comprobante(cp(moneda="USD", tipo_cambio="3.550"), CONTAB, MES, "080001")
    assert usd[0]["E"] == "US" and usd[0]["P"] == 100.0 and usd[0]["Q"] == "" and usd[2]["K"] == "421202"
    assert usd[0]["G"] == 3.55 and usd[0]["H"] == "C"            # el T.C. de SUNAT del comprobante, conversión especial
    sin_tc = driver_concar.filas_de_comprobante(cp(moneda="USD"), CONTAB, MES, "080001")
    assert sin_tc[0]["G"] == "" and sin_tc[0]["H"] == "V"        # sin T.C. → CONCAR lo busca en su tabla
    with pytest.raises(concar.SinCodigoDeMoneda) as e:
        driver_concar.filas_de_comprobante(cp(moneda="EUR"), CONTAB, MES, "080001")
    assert e.value.monedas == ["EUR"] and concar.monedas_sin_codigo([cp(moneda="EUR"), cp()], CONTAB) == ["EUR"]


def test_boleta_honorarios_nota_de_credito_y_tasa():
    # Boleta: no da crédito fiscal → todo al gasto, sin línea de IGV ni tasa, aunque traiga IGV
    bv = driver_concar.filas_de_comprobante(cp(tipo_cp="03", serie="B001", numero="55"), CONTAB, MES, "080001")
    assert len(bv) == 2 and bv[0]["B"] == "13" and bv[0]["R"] == "BV" and bv[0]["S"] == "B001-55"
    assert [f["O"] for f in bv] == [118.0, 118.0] and bv[0]["AO"] == "" and bv[1]["K"] == "421201"
    # Recibo por honorarios: sin IGV, cuenta por pagar 424101/424102, sin anexo auxiliar
    rh = driver_concar.filas_de_comprobante(cp(tipo_cp="02", serie="E001", numero="7", base_gravada="0", igv="0", inafecto="220", total="220"), CONTAB, MES, "080001")
    assert len(rh) == 2 and rh[0]["B"] == "15" and rh[0]["R"] == "RH" and rh[0]["AO"] == ""
    assert [f["O"] for f in rh] == [220.0, 220.0] and [f["N"] for f in rh] == ["D", "H"] and rh[1]["K"] == "424101" and rh[1]["X"] == ""
    assert driver_concar.filas_de_comprobante(cp(tipo_cp="02", moneda="USD", tipo_cambio="3.5"), CONTAB, MES, "080001")[1]["K"] == "424102"
    # Nota de crédito: invierte (proveedor al Debe, gasto e IGV al Haber) y lleva el documento que modifica
    nc = driver_concar.filas_de_comprobante(cp(tipo_cp="07", serie="FC01", numero="9", ref_tipo_cp="01", ref_serie="F001", ref_numero="00000123", ref_fecha="2026-08-01"),
                        CONTAB, MES, "080002")
    assert [f["N"] for f in nc] == ["H", "H", "D"] and nc[0]["R"] == "NC" and nc[0]["B"] == "11"
    assert (nc[0]["Z"], nc[0]["AA"], nc[0]["AB"]) == ("FT", "F001-123", date(2026, 8, 1))
    assert debe_haber(nc) == (Decimal("118"), Decimal("118"))
    nd = driver_concar.filas_de_comprobante(cp(tipo_cp="08", serie="FD01", numero="3"), CONTAB, MES, "080003")
    assert [f["N"] for f in nd] == ["D", "D", "H"]                # la nota de débito es una compra normal
    # Tasa IGV: la del comprobante, redondeada a entero (10.5 % → 11, John 10-sep-2026); sin IGV → vacía
    rest = driver_concar.filas_de_comprobante(cp(base_gravada="100", igv="10.5", total="110.5"), CONTAB, MES, "080004")
    assert rest[0]["AO"] == 11 and rest[1]["O"] == 10.5
    inaf = driver_concar.filas_de_comprobante(cp(base_gravada="0", igv="0", inafecto="118"), CONTAB, MES, "080005")
    assert len(inaf) == 2 and inaf[0]["AO"] == ""
    assert driver_concar.tasa_igv_entera(Decimal("18"), Decimal("0")) == ""      # IGV sin base: no hay tasa que leer, y no se inventa
    sin_venc = driver_concar.filas_de_comprobante(cp(fecha_vencimiento=None), CONTAB, MES, "080001")
    assert sin_venc[0]["U"] == EMISION                           # sin vencimiento → la de emisión
    sin_concepto = driver_concar.filas_de_comprobante(cp(concepto=""), CONTAB, MES, "080001")
    assert sin_concepto[0]["F"] == "ELECTROMECANICA DE PRUEBA S.R.L."   # glosa: si no hay concepto, el proveedor


def test_fecha_por_comprobante():
    """D y J (regla de un contador, 30-ago-2026, compras Y ventas): cada comprobante se asienta con
    SU fecha de emisión; el extemporáneo (mes anterior) cae al primer día del periodo y un
    emitido después del periodo, al último — en CONCAR el asiento cae en el mes de esa fecha
    y todo debe caer en el mes del proceso. T y U siguen siendo las del documento."""
    del_mes = driver_concar.filas_de_comprobante(cp(), CONTAB, MES, "080001")
    assert all(f["D"] == EMISION and f["J"] == EMISION for f in del_mes)
    junio = driver_concar.filas_de_comprobante(cp(fecha_emision="2026-06-15"), CONTAB, MES, "080002")
    assert all(f["D"] == date(2026, 8, 1) and f["J"] == date(2026, 8, 1) for f in junio)
    assert junio[0]["T"] == date(2026, 6, 15)                     # Fecha de Documento: SIEMPRE la original
    post = driver_concar.filas_de_comprobante(cp(fecha_emision="2026-09-02", fecha_vencimiento="2026-09-10"), CONTAB, MES, "080003")
    assert all(f["D"] == date(2026, 8, 31) and f["J"] == date(2026, 8, 31) for f in post)
    assert post[0]["T"] == date(2026, 9, 2) and post[0]["U"] == date(2026, 9, 10)
    sin_fecha = driver_concar.filas_de_comprobante(cp(fecha_emision=None, fecha_vencimiento=None), CONTAB, MES, "080004")
    assert all(f["D"] == date(2026, 8, 1) for f in sin_fecha)     # FECHA_FALTA: al primer día
    venta_jul = driver_concar.filas_de_comprobante(cp(fecha_emision="2026-07-20", cuenta_contable=""), CONTAB, MES, "080005", es_venta=True)
    assert all(f["D"] == date(2026, 8, 1) and f["J"] == date(2026, 8, 1) for f in venta_jul)


def test_recibo_por_honorarios_con_retencion_de_4ta():
    """El recibo que MUESTRA retención parte el asiento: gasto por el total, la retención
    en su cuenta y a la cuenta por pagar solo el neto. Sin retención (lo normal con
    suspensión) sigue siendo de dos líneas: nunca se calcula el 8 % solo."""
    rh = lambda **k: cp(tipo_cp="02", serie="E001", numero="7", base_gravada="0", igv="0",
                        inafecto="2000", total="2000", concepto="Asesoria contable", **k)
    con = driver_concar.filas_de_comprobante(rh(retencion="160"), CONTAB, MES, "080005")
    assert len(con) == 3 and [f["N"] for f in con] == ["D", "H", "H"]
    assert [f["K"] for f in con] == ["631101", "401721", "424101"]
    assert [f["O"] for f in con] == [2000.0, 160.0, 1840.0]        # gasto total · retención · NETO
    assert con[1]["W"] == "ASESORIA CONTABLE" and con[1]["L"] == "" and con[1]["M"] == ""
    assert con[2]["L"] == "20607777773"                             # el RUC va en la línea del profesional
    assert all(f["B"] == "15" and f["R"] == "RH" and f["AO"] == "" for f in con)
    assert debe_haber(con) == (Decimal("2000"), Decimal("2000"))
    # Sin retención: dos líneas, el total completo a la cuenta por pagar
    sin = driver_concar.filas_de_comprobante(rh(), CONTAB, MES, "080006")
    assert len(sin) == 2 and [f["O"] for f in sin] == [2000.0, 2000.0]
    # La cuenta de la retención se configura por RUC
    config = configuracion({"cuentas": {"retencion_4ta": "401722"}})
    assert driver_concar.filas_de_comprobante(rh(retencion="160"), config, MES, "080007")[1]["K"] == "401722"
    # En dólares, el neto va a 424102 y la retención también en USD
    usd = driver_concar.filas_de_comprobante(rh(retencion="16", moneda="USD", tipo_cambio="3.55"), CONTAB, MES, "080008")
    assert usd[1]["P"] == 16.0 and usd[2]["K"] == "424102" and usd[2]["P"] == 1984.0
    # Una retención en algo que no es recibo por honorarios NO parte el asiento
    assert len(driver_concar.filas_de_comprobante(cp(retencion="18"), CONTAB, MES, "080009")) == 3
    assert driver_concar.filas_de_comprobante(cp(retencion="18"), CONTAB, MES, "080009")[2]["O"] == 118.0


def test_el_interruptor_de_centros_de_costo():
    """Hay CONCARs que no llevan centros de costo: apagado, las columnas M (centro de
    costo) y X (anexo auxiliar) salen vacías aunque el comprobante traiga uno."""
    con = driver_concar.filas_de_comprobante(cp(), CONTAB, MES, "080001")
    assert con[0]["M"] == "OBRA01" and con[2]["X"] == "OBRA01"          # encendido (por defecto)
    config = configuracion({"usa_centros_costo": False})
    sin = driver_concar.filas_de_comprobante(cp(), config, MES, "080001")
    assert sin[0]["M"] == "" and sin[2]["X"] == ""
    assert [f["O"] for f in sin] == [100.0, 18.0, 118.0]                 # el asiento no cambia en nada más


def test_la_cuenta_decide_si_el_centro_va_a_la_M():
    """El contador, 09-sep-2026: «la cuenta 63 y 65 tiene habilitado el centro de costo en la
    columna M, pero cuando es una cuenta 60 por defecto no se debe asignar un centro de costo».
    En CONCAR esa marca vive en cada cuenta del plan; aquí, en `cuentas_con_centro` por prefijo.

    OJO al leer este test: el fixture usa 631101, que SÍ está en la lista de fábrica. Por eso
    ningún test anterior se enteró del cambio, y por eso hace falta este.
    """
    # Las de la lista de fábrica: 62, 63, 65 y —para que ventas no cambie— 70.
    assert driver_concar.filas_de_comprobante(cp(), CONTAB, MES, "080001")[0]["M"] == "OBRA01"                      # 631101
    assert driver_concar.filas_de_comprobante(cp(cuenta_contable="659999"), CONTAB, MES, "080001")[0]["M"] == "OBRA01"
    # El 62 entró el 22-sep-2026 (John), para cualquier sistema: una 62 es gasto y lleva su centro como las otras
    # tres. Hasta ese día esta misma cuenta estaba abajo, entre las que NO lo llevan.
    assert driver_concar.filas_de_comprobante(cp(cuenta_contable="627401"), CONTAB, MES, "080001")[0]["M"] == "OBRA01"
    # Las que NO: ni M ni X. Son las clases reales de un constructor: 603 compras, 681 depreciación.
    for cuenta in ("601101", "603201", "681401"):
        gasto = driver_concar.filas_de_comprobante(cp(cuenta_contable=cuenta), CONTAB, MES, "080001")[0]
        assert (gasto["K"], gasto["M"], gasto["X"]) == (cuenta, "", ""), cuenta
    # Y el doble anexo del proveedor NO depende de esto: sigue llevando el centro.
    assert driver_concar.filas_de_comprobante(cp(cuenta_contable="603201"), CONTAB, MES, "080001")[2]["X"] == "OBRA01"
    # Ventas sin cambios: 701101 está en la lista, que es exactamente para lo que se metió el 70. Es la
    # TERCERA fila desde la 2.2: el orden de la venta es cliente · IGV · ingreso.
    filas_venta = driver_concar.filas_de_comprobante(cp(cuenta_contable=""), CONTAB, MES, "080001", es_venta=True)
    assert filas_venta[2]["K"] == "701101" and filas_venta[2]["M"] == "OBRA01"
    # Casa por prefijo, no por los dos primeros dígitos: "6311" no alcanza a 631201.
    largo = configuracion({"cuentas_con_centro": ["6311"]})
    assert driver_concar.filas_de_comprobante(cp(), largo, MES, "080001")[0]["M"] == "OBRA01"                        # 631101
    assert driver_concar.filas_de_comprobante(cp(cuenta_contable="631201"), largo, MES, "080001")[0]["M"] == ""
    # Lista VACÍA es una respuesta legítima —ninguna cuenta lo lleva— y no es lo mismo que ausente.
    vacia = configuracion({"cuentas_con_centro": []})
    assert driver_concar.filas_de_comprobante(cp(), vacia, MES, "080001")[0]["M"] == ""
    assert concar.lleva_centro("631101", {}) is True                                             # sin la clave: los de fábrica


def test_el_centro_de_referencia_en_la_X_del_gasto():
    """«Algunas empresas optan en colocar la columna X como referencia el centro de costo»
    (el contador, 09-sep-2026). Sin marcar de fábrica: la X de CONCAR solo admite dato si esa
    cuenta tiene anexo referencia, y escribirla donde no toca puede tumbar la importación."""
    ref = configuracion(REFERENCIA_EN_X)
    gasto, _, prov = driver_concar.filas_de_comprobante(cp(cuenta_contable="603201"), ref, MES, "080001")
    assert (gasto["M"], gasto["X"]) == ("", "OBRA01")          # la referencia, en su propia línea
    assert prov["X"] == "OBRA01"                                # y el doble anexo, intacto a la vez
    # A las cuentas que SÍ lo llevan en M, el interruptor no las toca.
    con_m = driver_concar.filas_de_comprobante(cp(), ref, MES, "080001")[0]
    assert (con_m["M"], con_m["X"]) == ("OBRA01", "")
    # El interruptor maestro manda sobre los dos.
    apagado = configuracion({**REFERENCIA_EN_X, "usa_centros_costo": False})
    sin = driver_concar.filas_de_comprobante(cp(cuenta_contable="603201"), apagado, MES, "080001")[0]
    assert (sin["M"], sin["X"]) == ("", "")
    # La línea de la detracción sigue limpia aunque la referencia esté encendida.
    det = driver_concar.filas_de_comprobante(cp(cuenta_contable="603201", detraccion={"codigo": "027", "porcentaje": 4}), ref, MES, "100001")
    assert (det[-1]["R"], det[-1]["M"], det[-1]["X"]) == ("DR", "", "")


def test_las_columnas_elegidas_deciden_en_que_anexo_va_el_centro():
    """El centro se guarda una vez y la configuración de CONCAR elige sus columnas (13-sep-2026): la M va siempre; la
    X del tercero, marcada de fábrica, se puede quitar, y la X del gasto se puede elegir sin la del tercero."""
    from contaperu import drivers
    from contaperu.drivers import contrato

    assert contrato.centro_en_anexo(driver_concar, CONTAB) == {"tercero"}
    solo_m = configuracion({"columnas": {"centro_costo": ["centro_costo"]}})
    assert contrato.centro_en_anexo(driver_concar, solo_m) == frozenset()
    gasto, _, proveedor = driver_concar.filas_de_comprobante(cp(), solo_m, MES, "080001")
    assert (gasto["M"], gasto["X"], proveedor["X"]) == ("OBRA01", "", "")
    solo_referencia = configuracion({"columnas": {"centro_costo": ["centro_costo", "anexo_auxiliar"]}})
    gasto, _, proveedor = driver_concar.filas_de_comprobante(cp(cuenta_contable="603201"), solo_referencia, MES,
                                                             "080001")
    assert (gasto["M"], gasto["X"], proveedor["X"]) == ("", "OBRA01", "")
    # Un driver que no declara columnas lleva lo del estándar: el centro en el anexo del tercero.
    assert contrato.centro_en_anexo(drivers.csv, solo_m) == {"tercero"}


def test_el_centro_solo_es_obligatorio_donde_se_escribe():
    """Una cuenta que no lleva centro no puede bloquear la exportación de un mes."""
    assert concar.comprobantes_sin_centro([cp(cuenta_contable="603201", centro_costo="")], CONTAB) == []
    assert [c.numero for c in concar.comprobantes_sin_centro([cp(centro_costo="")], CONTAB)] == ["00000123"]
    # Y una 62 sí lo exige desde el 22-sep-2026, porque desde entonces lo escribe: las dos caras de la misma
    # lista. Un mes con cuentas 62 y sin centro deja de estar listo, y eso es lo que se quiere.
    assert [c.numero for c in concar.comprobantes_sin_centro(
        [cp(cuenta_contable="627401", centro_costo="")], CONTAB)] == ["00000123"]
    # Ni siquiera con la referencia en X encendida: esa X es una referencia, y una referencia
    # que se puede dejar en blanco no puede impedir exportar.
    ref = configuracion(REFERENCIA_EN_X)
    assert concar.comprobantes_sin_centro([cp(cuenta_contable="603201", centro_costo="")], ref) == []
    # Ventas: sin el flag la cuenta se resolveria como gasto. Con `cuentas.gasto` vacío eso da cuenta vacía → no
    # bloquea; con el flag cae en 701101 → sí bloquea. Va con SIN_GASTO porque con la de CONCAR (631101) la cuenta
    # ya no sale vacía y el caso que se quiere probar no existiría.
    vacio = cp(cuenta_contable="", centro_costo="")
    assert concar.comprobantes_sin_centro([vacio], SIN_GASTO) == []
    assert [c.numero for c in concar.comprobantes_sin_centro([vacio], SIN_GASTO, es_venta=True)] == ["00000123"]


def test_factura_con_detraccion_va_al_sub_diario_10():
    """La factura con detracción, calcada dun Excel real de produccion (06-sep-2026): el total COMPLETO al
    proveedor (421201) y dos líneas más por el monto detraído — el proveedor al Debe y 421203 al Haber
    con tipo DT, comodín 9999999999 y AI–AL. El monto va en soles enteros.

    La glosa es la misma que la de las otras cuatro filas desde la 2.2: la línea se reconoce por su cuenta y su
    tipo de documento, no por un texto."""
    det = {"codigo": "037", "porcentaje": "12", "monto": "14.16", "cuenta": "00-000-123456"}
    c = cp(detraccion=det)
    assert concar.tiene_detraccion(c) and concar.sub_diario(c, CONTAB) == "10" and concar.sigla_documento(c, CONTAB) == "FT"
    filas = driver_concar.filas_de_comprobante(c, CONTAB, MES, "080001")
    assert len(filas) == 5 and all(f["B"] == "10" for f in filas)
    gasto, igv, cxp, det_prov, det_ = filas
    assert (cxp["K"], cxp["N"], cxp["O"], cxp["AI"]) == ("421201", "H", 118.0, "")      # el total completo, sin AI
    assert (det_prov["K"], det_prov["L"], det_prov["N"], det_prov["O"], det_prov["R"], det_prov["S"], det_prov["X"]) == \
        ("421201", "20607777773", "D", 14.0, "FT", "F001-123", "OBRA01")
    assert (det_["K"], det_["L"], det_["N"], det_["O"], det_["Q"]) == ("421203", "20607777773", "H", 14.0, 14.0)   # 118 × 12 % = 14.16 → 14
    assert (det_["R"], det_["S"], det_["M"], det_["X"]) == ("DR", "9999999999", "", "")
    assert det_["W"] == gasto["W"], "la misma glosa en las cinco filas"
    assert (det_["AI"], det_["AJ"], det_["AK"], det_["AL"], det_["AO"]) == ("03701", 12.0, "", 118.0, 18)
    assert (gasto["AI"], igv["AI"], cxp["AI"]) == ("", "", "") and debe_haber(filas) == (Decimal("132"), Decimal("132"))
    # Dólares: la detracción se deposita en SOLES → 118 × 3.5 = 413 × 12 % = 49.56 → 50 soles → 14.29 US en la línea
    usd = driver_concar.filas_de_comprobante(cp(detraccion=det, moneda="USD", tipo_cambio="3.5"), CONTAB, MES, "080001")
    assert len(usd) == 5 and usd[2]["K"] == "421202" and usd[3]["K"] == "421202"
    assert (usd[4]["K"], usd[4]["O"], usd[4]["P"], usd[4]["Q"], usd[4]["AK"], usd[4]["AL"]) == ("421203", 14.29, 14.29, "", 118.0, "")
    assert len(driver_concar.filas_de_comprobante(cp(detraccion=det, moneda="USD", tipo_cambio=None), CONTAB, MES, "080001")) == 3   # sin T.C. no se puede calcular
    propia = configuracion({"cuentas": {"cxp_detraccion": {"PEN": "421209"}}})
    assert driver_concar.filas_de_comprobante(c, propia, MES, "080001")[4]["K"] == "421209"       # la cuenta sale de la tabla de Configuración
    # Elegida en pantalla (05-sep-2026): el desplegable de Revisión guarda {codigo, porcentaje de la tabla,
    # monto informativo, cuenta ""} y el asiento es el mismo que con la detracción del XML.
    pantalla = cp(detraccion={"codigo": "037", "porcentaje": "12", "monto": "14.16", "cuenta": ""})
    fp = driver_concar.filas_de_comprobante(pantalla, CONTAB, MES, "080001")
    assert [f["B"] for f in fp] == ["10"] * 5 and (fp[4]["K"], fp[4]["AI"], fp[4]["AJ"]) == ("421203", "03701", 12.0)
    # Código fuera de la tabla y sin tasa: no se sabe cuánto detraer → el asiento normal de 3 líneas
    otro = driver_concar.filas_de_comprobante(cp(detraccion={"codigo": "031", "porcentaje": "", "monto": "0"}), CONTAB, MES, "080001")
    assert len(otro) == 3 and otro[2]["K"] == "421201" and otro[2]["AI"] == ""
    en_tabla = driver_concar.filas_de_comprobante(cp(detraccion={"codigo": "012", "porcentaje": "", "monto": "0"}), CONTAB, MES, "080001")[4]
    assert (en_tabla["AI"], en_tabla["AJ"]) == ("01201", 12.0)
    sin_tasa = driver_concar.filas_de_comprobante(cp(detraccion={"codigo": "019"}), CONTAB, MES, "080001")[4]
    assert (sin_tasa["AI"], sin_tasa["AJ"], sin_tasa["O"]) == ("01903", 10.0, 12.0)     # 118 × 10 % = 11.8 → 12
    # El sub-diario NO manda: la empresa que lo lleva todo en el 11 hace el mismo asiento de 5 líneas
    config = configuracion({"sub_diario_detraccion": "", "detraccion_codigos": {"037": "03799"}})
    en11 = driver_concar.filas_de_comprobante(c, config, MES, "080001")
    assert concar.sub_diario(c, config) == "11" and [f["B"] for f in en11] == ["11"] * 5 and en11[4]["AI"] == "03799"
    # El recibo por honorarios nunca lleva detracción
    rh = cp(tipo_cp="02", detraccion=det)
    assert concar.sub_diario(rh, CONTAB) == "15" and all(f["AI"] == "" and f["R"] != "DR" for f in driver_concar.filas_de_comprobante(rh, CONTAB, MES, "080001"))
    assert not concar.tiene_detraccion(cp(detraccion={"codigo": "", "porcentaje": "0"})) and not concar.tiene_detraccion(cp(detraccion=None))


def test_el_codigo_de_area_es_de_la_empresa_y_solo_va_en_la_detraccion():
    """La columna V (Tabla General 26 de CONCAR) sale del Excel que un CONCAR real aceptó.

    Es un número PROPIO de cada empresa, no una constante contable, así que sale vacío de fábrica:
    ponerle uno por defecto metería los apuntes de todo el mundo en un área que nadie eligió. Y va
    en la línea de la detracción y en ninguna otra — en el archivo validado las otras cuatro la
    tenían en blanco.
    """
    c = cp(detraccion={"codigo": "027", "porcentaje": "4"})

    # Sin configurar: nada cambia para quien ya exportaba.
    assert [f["V"] for f in driver_concar.filas_de_comprobante(c, CONTAB, MES, "080001")] == [""] * 5

    con_area = configuracion({"detraccion_area": "900"})
    filas = driver_concar.filas_de_comprobante(c, con_area, MES, "080001")
    assert [f["V"] for f in filas] == ["", "", "", "", "900"]
    # Y no se cuela en el resto del asiento ni en un comprobante sin detracción.
    sin_det = driver_concar.filas_de_comprobante(cp(), con_area, MES, "080001")
    assert len(sin_det) == 3 and [f["V"] for f in sin_det] == [""] * 3


def test_el_codigo_de_area_no_se_recorta():
    """La plantilla dice «3 Caracteres», pero recortar un CÓDIGO es peor que pasarse.

    `9001` recortado a `900` manda el apunte a OTRA área **en silencio** y nadie se entera hasta que
    cuadran el área a fin de mes. Entero, CONCAR lo rechaza en la importación y se ve al momento.
    Es la diferencia con la glosa, que sí se corta: ahí sobra texto, aquí sobraría significado.
    Y desde el 13-sep-2026 ni llega al Excel: la configuración de CONCAR declara su largo y lo rechaza antes.
    """
    c = cp(detraccion={"codigo": "027", "porcentaje": "4"})
    filas = driver_concar.filas_de_comprobante(c, dict(CONTAB, detraccion_area="9001"), MES, "080001")
    assert filas[-1]["V"] == "9001"
    with pytest.raises(api.ConfiguracionInvalida) as e:
        configuracion({"detraccion_area": "9001"})
    assert e.value.errores == ['`concar.detraccion_area`: el texto "9001" no cumple el patrón ^[A-Z0-9]{0,3}$']


def test_el_tipo_de_documento_de_la_detraccion_es_configurable():
    """`DR` es lo que aceptó un CONCAR real, pero la Tabla General 06 la numera cada contribuyente."""
    c = cp(detraccion={"codigo": "027", "porcentaje": "4"})
    assert driver_concar.filas_de_comprobante(c, CONTAB, MES, "080001")[-1]["R"] == "DR"
    otro = configuracion({"detraccion_tipo_doc": "DT"})
    assert driver_concar.filas_de_comprobante(c, otro, MES, "080001")[-1]["R"] == "DT"


def test_la_detraccion_referencia_al_documento_del_que_sale():
    """Z/AA/AB en la línea de la detracción: de qué documento sale este depósito.

    En una factura ese documento es el propio comprobante, y es lo que CONCAR aceptó. Las otras
    cuatro líneas siguen sin referencia: en el archivo validado estaban en blanco.
    """
    filas = driver_concar.filas_de_comprobante(cp(detraccion={"codigo": "027", "porcentaje": "4"}), CONTAB, MES, "080001")
    assert [f["Z"] for f in filas] == ["", "", "", "", "FT"]
    assert [f["AA"] for f in filas] == ["", "", "", "", "F001-123"]
    assert [f["AB"] for f in filas] == ["", "", "", "", date(2026, 8, 11)]


def test_una_nota_con_detraccion_conserva_la_referencia_a_la_factura():
    """El caso que la regla nueva podía pisar sin querer.

    Una nota de crédito ya llevaba en Z/AA/AB la factura que corrige, en TODAS sus líneas. La
    detracción no se la quita: no hay ningún archivo validado que diga que una nota deba referenciar
    a sí misma, y cambiarlo sería inventarse una regla contable. Queda pendiente de comprobar con
    una nota de crédito real que entre en CONCAR.
    """
    nc = cp(tipo_cp="07", serie="FC01", numero="9", detraccion={"codigo": "027", "porcentaje": "4"},
            ref_tipo_cp="01", ref_serie="F001", ref_numero="123", ref_fecha="2026-08-10")
    filas = driver_concar.filas_de_comprobante(nc, CONTAB, MES, "080001")
    assert len(filas) == 5 and filas[-1]["R"] == "DR"
    # La referencia es la FACTURA (F001-123), no la propia nota (FC01-9), en las cinco líneas.
    assert [f["AA"] for f in filas] == ["F001-123"] * 5
    assert filas[-1]["Z"] == "FT" and filas[-1]["AB"] == date(2026, 8, 10)


def test_asiento_con_detraccion_calca_un_excel_real():
    """La factura E001-871 dun Excel real que un CONCAR de produccion acepto en CONCAR (06-sep-2026): las cinco filas, celda a celda."""
    c = cp(serie="E001", numero="871", fecha_emision="2026-08-10", fecha_vencimiento="2026-08-27",
           contraparte_doc="20602222226", contraparte_nombre="PROVEEDOR DE PRUEBA SAC",
           base_gravada="4200", igv="756", total="4956", concepto="SERVICIO DE TRANSPORTE DE MATERIALES BENCE",
           cuenta_contable="659999", centro_costo="CC-64", detraccion={"codigo": "027", "porcentaje": "4"})
    filas = driver_concar.filas_de_comprobante(c, CONTAB, MES, "080084")
    col = lambda k: [f[k] for f in filas]
    assert len(filas) == 5
    assert col("B") == ["10"] * 5 and col("C") == ["080084"] * 5 and col("E") == ["MN"] * 5 and col("H") == ["V"] * 5 and col("I") == ["S"] * 5
    assert col("F") == ["SERVICIO DE TRANSPORTE DE MATERIALES BEN"] * 5
    assert col("K") == ["659999", "401111", "421201", "421201", "421203"]
    assert col("L") == ["", "", "20602222226", "20602222226", "20602222226"]
    assert col("M") == ["CC-64", "", "", "", ""] and col("X") == ["", "", "CC-64", "CC-64", ""]
    assert col("N") == ["D", "D", "H", "D", "H"]
    assert col("O") == [4200.0, 756.0, 4956.0, 198.0, 198.0] and col("Q") == col("O") and col("P") == [""] * 5
    assert col("R") == ["FT", "FT", "FT", "FT", "DR"] and col("S") == ["E001-871"] * 4 + ["9999999999"]
    assert col("D") == [date(2026, 8, 10)] * 5 and col("T") == [date(2026, 8, 10)] * 5 and col("U") == [date(2026, 8, 27)] * 5
    assert col("W") == ["SERVICIO DE TRANSPORTE DE MATE"] * 5, "la misma glosa en las cinco filas (2.2)"
    assert col("AI") == ["", "", "", "", "02702"] and col("AJ") == ["", "", "", "", 4.0]
    assert col("AL") == ["", "", "", "", 4956.0] and col("AK") == [""] * 5
    assert col("AO") == [18] * 5 and debe_haber(filas) == (Decimal("5154"), Decimal("5154"))
    # La referencia: solo la línea de la detracción dice de qué documento sale, y ese documento es
    # la propia factura. Las otras cuatro no la llevan — así estaba en el archivo validado.
    assert col("Z") == ["", "", "", "", "FT"] and col("AA") == ["", "", "", "", "E001-871"]
    assert col("AB") == ["", "", "", "", date(2026, 8, 10)]
    # Y el área sale VACÍA mientras el contribuyente no ponga la suya: `CONTAB` no la configura.
    assert col("V") == [""] * 5


def test_asiento_de_ventas_espejo_del_skill():
    """Ventas: cliente 121201/121202 al Debe por el total, IGV 401111 al Haber, ingreso 701101 al
    Haber por el valor venta; sub-diario 05; NC invierte; cada venta con su fecha de emisión; la boleta de
    venta SÍ lleva su IGV; la detracción no se registra.

    El IGV va ANTES del ingreso desde la 2.2: es el orden de los archivos reales de STARSOFT, y el que el
    motor ya usaba en compras."""
    fv = driver_concar.filas_de_comprobante(cp(cuenta_contable=""), CONTAB, MES, "080009", es_venta=True)   # sin cuenta en la fila → default 701101
    cli_, igv, ing = fv
    assert [f["N"] for f in fv] == ["D", "H", "H"] and [f["K"] for f in fv] == ["121201", "401111", "701101"]
    assert [f["O"] for f in fv] == [118.0, 18.0, 100.0]
    assert cli_["L"] == "20607777773" and ing["L"] == "" and cli_["X"] == "OBRA01"    # anexo RUC + anexo auxiliar en el cliente
    assert ing["M"] == "OBRA01" and cli_["M"] == ""                                   # centro de costo en el ingreso
    assert all(f["B"] == "05" and f["D"] == EMISION and f["J"] == EMISION for f in fv)  # cada venta en su fecha
    assert fv[0]["AO"] == 18 and fv[0]["R"] == "FT"
    assert debe_haber(fv) == (Decimal("118"), Decimal("118"))
    # La cuenta de ingreso: default real 701101; la fila o el RUC pueden cambiarla
    assert driver_concar.filas_de_comprobante(cp(cuenta_contable="702101"), CONTAB, MES, "080001", es_venta=True)[2]["K"] == "702101"
    config = configuracion({"cuentas": {"ventas": "701201", "clientes": {"USD": "121209"}}})
    v2 = driver_concar.filas_de_comprobante(cp(cuenta_contable="", moneda="USD", tipo_cambio="3.55"), config, MES, "080001", es_venta=True)
    assert v2[2]["K"] == "701201" and v2[0]["K"] == "121209" and v2[0]["E"] == "US" and v2[0]["G"] == 3.55
    assert concar.comprobantes_sin_cuenta([cp(cuenta_contable="")], CONTAB, es_venta=True) == []
    # Boleta de venta emitida: SÍ lleva su IGV (la regla "sin crédito" es solo de compras)
    bv = driver_concar.filas_de_comprobante(cp(tipo_cp="03", serie="B001", numero="9", cuenta_contable=""), CONTAB, MES, "080002", es_venta=True)
    assert len(bv) == 3 and bv[0]["B"] == "05" and bv[0]["R"] == "BV" and bv[1]["K"] == "401111" and bv[0]["AO"] == 18
    # NC de venta: invierte (ingreso D, IGV D, cliente H) con su documento de referencia
    nc = driver_concar.filas_de_comprobante(cp(tipo_cp="07", serie="FC01", numero="3", cuenta_contable="", ref_tipo_cp="01", ref_serie="F001", ref_numero="00000123", ref_fecha="2026-08-01"),
                        CONTAB, MES, "080003", es_venta=True)
    assert [f["N"] for f in nc] == ["D", "D", "H"] and [f["K"] for f in nc] == ["701101", "401111", "121201"]
    assert nc[0]["Z"] == "FT" and nc[0]["AA"] == "F001-123" and debe_haber(nc) == (Decimal("118"), Decimal("118"))
    # La detracción en ventas no se registra: sigue en el 05 y sin columnas AI-AL
    det = cp(detraccion={"codigo": "037", "porcentaje": "12", "monto": "14.16"})
    assert concar.sub_diario(det, CONTAB, es_venta=True) == "05"
    assert driver_concar.filas_de_comprobante(det, CONTAB, MES, "080004", es_venta=True)[0]["AI"] == ""
    # Numeración de ventas: un solo sub-diario 05
    cs = [cp(numero="1"), cp(tipo_cp="03", serie="B001", numero="2")]
    numeros, rangos = concar.numerar(cs, CONTAB, "202608", {"05": 9}, es_venta=True)
    assert [numeros[id(c)] for c in cs] == ["080009", "080010"] and rangos["05"]["hasta"] == 10


def test_el_codigo_sunat_manda_y_lo_de_concar_se_deriva():
    """Un tipo SUNAT sin sigla/sub-diario configurado no se inventa; por RUC se puede añadir o cambiar."""
    assert concar.sigla_documento(cp(tipo_cp="02"), CONTAB) == "RH" and concar.sub_diario(cp(tipo_cp="02"), CONTAB) == "15"
    luz = cp(tipo_cp="14", serie="", numero="12345")             # recibo de luz: SUNAT 14 → RC en el 11 (el contador, 23-ago-2026)
    assert concar.sigla_documento(luz, CONTAB) == "RC" and concar.sub_diario(luz, CONTAB) == "11"
    assert driver_concar.filas_de_comprobante(luz, CONTAB, MES, "080001")[0]["S"] == "12345" and len(driver_concar.filas_de_comprobante(luz, CONTAB, MES, "080001")) == 3
    bancario = cp(tipo_cp="13", serie="", numero="77")           # documento bancario: sin entrada por defecto
    assert concar.sigla_documento(bancario, CONTAB) == "" and concar.tipos_sin_sigla([bancario, cp()], CONTAB) == ["13"]
    with pytest.raises(concar.SinSigla) as e:
        construir_concar(COMPRAS, [bancario], CONTAB, {"11": 1})
    assert e.value.tipos == ["13"]
    config = configuracion({"tipos": {"13": {"sigla": "DB", "sub_diario": "11"}, "01": {"sigla": "FA", "sub_diario": "12"}}})
    assert concar.tipos_sin_sigla([bancario], config) == []
    assert driver_concar.filas_de_comprobante(bancario, config, MES, "080001")[0]["R"] == "DB"
    assert driver_concar.filas_de_comprobante(cp(), config, MES, "080001")[0]["B"] == "12" and driver_concar.filas_de_comprobante(cp(), config, MES, "080001")[0]["R"] == "FA"
    assert concar.sub_diario(cp(tipo_cp="03"), config) == "13"   # lo no tocado conserva el default


def test_la_configuracion_del_entorno_se_funde_con_los_valores_por_defecto():
    """Lo general en la raíz y lo de CONCAR en su sección (John, 13-sep-2026): lo que el entorno cambia manda, y lo que
    no toca conserva su valor por defecto, también dentro de un mismo grupo (`fundir_config` funde en profundidad)."""
    entorno = {"cuentas": {"gasto": "659301", "cxp": {"USD": "421203"}},
               "concar": {"tipos": {"20": {"sigla": "CO", "sub_diario": "11"}}}}
    c = con_imputaciones(prep.config_aplicada(entorno, "concar"))
    assert c["cuentas"]["gasto"] == "659301"                              # el entorno manda
    assert c["cuentas"]["cxp"] == {"PEN": "421201", "USD": "421203"}      # se funde dentro del grupo
    assert c["cuentas"]["igv"] == "401111" and c["tipos"]["01"]["sigla"] == "FT"
    assert c["tipos"]["20"]["sigla"] == "CO"                              # un tipo añadido por el entorno
    assert configuracion(None) == configuracion({}) == CONTAB


def test_cuenta_obligatoria_y_default_del_ruc():
    # Sin ninguna cuenta de gasto no hay de dónde sacarla, y se planta: es el caso de quien la deja en blanco a
    # propósito para que le avisen.
    with pytest.raises(concar.SinCuenta):
        driver_concar.filas_de_comprobante(cp(cuenta_contable=""), SIN_GASTO, MES, "080001")
    # Con la que trae CONCAR, esa misma fila sale a su cuenta de gasto en vez de detener la exportación.
    sale = driver_concar.filas_de_comprobante(cp(cuenta_contable=""), CONTAB, MES, "080001")
    assert sale[0]["K"] == "631101"
    config = configuracion({"cuentas": {"gasto": "659901", "cxp": {"USD": "421203"}}})
    filas = driver_concar.filas_de_comprobante(cp(cuenta_contable="", moneda="USD"), config, MES, "080001")
    assert filas[0]["K"] == "659901" and filas[2]["K"] == "421203"
    assert config["cuentas"]["cxp"]["PEN"] == "421201" and config["cuentas"]["igv"] == "401111"   # lo no tocado se conserva
    assert config["cuentas"]["honorarios"] == {"PEN": "424101", "USD": "424102"}
    assert concar.comprobantes_sin_cuenta([cp(cuenta_contable=""), cp()], SIN_GASTO)[0].numero == "00000123"
    assert concar.comprobantes_sin_cuenta([cp(cuenta_contable="")], config) == []
    # Y con la de CONCAR tampoco falta ninguna: es la consecuencia de que el sistema traiga cuenta de gasto —el
    # mes sale «listo» y esas filas se imputan ahí—, y quien prefiera que le avisen deja la cuenta en blanco.
    assert concar.comprobantes_sin_cuenta([cp(cuenta_contable="")], CONTAB) == []
    # El centro de costo, igual (06-sep-2026): obligatorio con los centros encendidos; nada si están apagados
    assert [c.numero for c in concar.comprobantes_sin_centro([cp(centro_costo=""), cp(centro_costo="  "), cp()], CONTAB)] == ["00000123", "00000123"]
    assert concar.comprobantes_sin_centro([cp(centro_costo="")], configuracion({"usa_centros_costo": False})) == []


# ── Numeración MMNNNN por sub-diario ─────────────────────────────────────────

def test_numerar_por_sub_diario_con_el_mes_del_periodo():
    cs = [cp(numero="1"), cp(tipo_cp="03", serie="B001", numero="2"), cp(numero="3"), cp(tipo_cp="02", serie="E001", numero="4", igv="0"),
          cp(numero="5", detraccion={"codigo": "037", "porcentaje": "12"})]
    numeros, rangos = concar.numerar(cs, CONTAB, "202608", {"11": 20, "13": 1, "15": 5, "10": 8})
    assert [numeros[id(c)] for c in cs] == ["080020", "080001", "080021", "080005", "080008"]
    assert rangos["11"] == {"desde": 20, "hasta": 21, "comprobantes": 2, "desde_codigo": "080020", "hasta_codigo": "080021", "desborda": False}
    assert rangos["13"]["hasta"] == 1 and rangos["15"]["hasta"] == 5 and rangos["10"]["hasta"] == 8
    with pytest.raises(concar.SinCorrelativo) as e:
        concar.numerar(cs, CONTAB, "202608", {"11": 20})
    assert e.value.sub_diarios == ["13", "15", "10"]
    _, r = concar.numerar(cs[:1], CONTAB, "202612", {"11": 9999})
    assert r["11"]["desde_codigo"] == "129999"
    _, r = concar.numerar(cs[:3], CONTAB, "202612", {"11": 9999, "13": 1})
    assert r["11"]["desborda"] is True


# ── El .xlsx ─────────────────────────────────────────────────────────────────

def test_xlsx_con_la_plantilla_de_concar():
    xlsx, resumen = construir_concar(COMPRAS, [cp(), cp(tipo_cp="03", serie="B001", numero="55", igv="0", base_gravada="0", inafecto="118")],
                                     CONTAB, {"11": 20, "13": 1})
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb["CONCAR"]
    encabezados = [ws.cell(row=1, column=i).value for i in range(1, 42)]
    assert encabezados == list(driver_concar.datos.CABECERAS["titulos"].values())
    assert encabezados[0] == "WE" and encabezados[40] == "Tasa IGV" and ws.cell(row=1, column=42).value is None
    assert ws["C2"].value.startswith("Los dos primeros dígitos son el mes") and ws["B3"].value == "4 Caracteres"
    # Formato de la plantilla de un contador (30-ago-2026): titulo azul marino con letra
    # blanca bold, notas SIN relleno, alturas 45/120/30 (+15.8 por fila de datos),
    # panel congelado en A4 y autofiltro sobre la fila de formatos.
    assert ws["A1"].fill.start_color.rgb.endswith("191970") and ws["A1"].font.bold and ws["A1"].font.color.rgb.endswith("FFFFFF")
    assert ws["A2"].fill.patternType is None and ws["A3"].font.bold
    assert (ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height) == (45, 120, 30)
    assert ws.row_dimensions[4].height == 15.8 and ws.freeze_panes == "A4" and ws.auto_filter.ref == "A3:AO3"
    # OJO: acceder a column_dimensions["AO"] la CREA con el default de openpyxl (13);
    # el "sin ancho propio" se pregunta por membresia, no por el valor.
    assert "AO" not in ws.column_dimensions and ws.column_dimensions["F"].width == 41.43
    # Datos desde la fila 4: factura (3 filas) + boleta sin IGV (2 filas)
    assert [ws[f"N{r}"].value for r in range(4, 9)] == ["D", "D", "H", "D", "H"]
    assert [ws[f"C{r}"].value for r in range(4, 9)] == ["080020"] * 3 + ["080001"] * 2
    assert ws["C4"].number_format == "@" and ws["K5"].number_format == "@"                       # texto: CONCAR lee caracteres
    assert ws["O4"].value == 100.0 and ws["O4"].number_format == "#,##0.00" and ws["K5"].value == "401111"
    assert ws["D4"].value.date() == EMISION and ws["D4"].number_format == "dd/mm/yyyy"            # fecha real de Excel
    assert ws["T4"].value.date() == EMISION and ws["U4"].value.date() == VENCE and ws["J4"].value.date() == EMISION
    assert ws["L6"].value == "20607777773" and ws["M4"].value == "OBRA01" and ws["X6"].value == "OBRA01" and ws["S7"].value == "B001-55"
    assert ws["B7"].value == "13" and ws["R7"].value == "BV" and ws["AO7"].value is None and ws["AO4"].value == 18
    assert ws["O4"].font.name == "Aptos Narrow"
    assert resumen["filas"] == 5 and resumen["debe"] == resumen["haber"] == "236.00"
    assert resumen["sub_diarios"]["11"]["hasta_codigo"] == "080020" and resumen["sub_diarios"]["13"]["etiqueta"] == "Boletas de venta"
    assert resumen["fechas"] == "por comprobante (extemporáneos al 01/08/2026)"


def test_generar_con_plantilla_concar():
    exp = gen.generar(COMPRAS, [cp(), cp(numero="2", excluida=True)], "concar", config=CONTAB, correlativos={"11": 1})
    assert exp.nombre == exp.archivo == "CONCAR_COMPRAS_202608_20601111111.xlsx"
    assert exp.formato == "concar_xlsx" and exp.content_type.endswith("spreadsheetml.sheet")
    assert exp.contenido[:2] == b"PK" and exp.comprimido == b"" and exp.nombre_comprimido == ""
    assert exp.comprobantes == 1 and exp.resumen["comprobantes"] == 1 and exp.resumen["excluidos"] == 1 and exp.resumen["filas"] == 3
    expv = gen.generar(VENTAS, [cp()], "concar", config=CONTAB, correlativos={"05": 1})
    assert expv.nombre == "CONCAR_VENTAS_202608_20601111111.xlsx" and expv.resumen["sub_diarios"]["05"]["etiqueta"] == "Ventas"
    with pytest.raises(concar.SinCuenta):
        gen.generar(COMPRAS, [cp(cuenta_contable="")], "concar", config=SIN_GASTO, correlativos={"11": 1})
    with pytest.raises(concar.SinCodigoDeMoneda):
        gen.generar(COMPRAS, [cp(moneda="EUR")], "concar", config=CONTAB, correlativos={"11": 1})


# ── Endpoint ──────────────────────────────────────────────────────────────────

def _fila(pro, **k):
    base = dict(cp(**k).a_dict(), id=str(uuid.uuid4()), proceso_id=pro["id"], empresa_id=pro["empresa_id"],
                cliente_id=pro["cliente_id"], origen="manual", estado="ok", excluida=False, observaciones=[])
    return base


def test_sub_diario_general_de_compras_gobierna_los_tipos_sin_registro_propio():
    # Factura, ticket y nota caen al general; renumerarlo los mueve a TODOS a la vez
    config = configuracion({"sub_diario_compras": "20"})
    assert concar.sub_diario(cp(), config) == "20"
    assert concar.sub_diario(cp(tipo_cp="12"), config) == "20"
    assert concar.sub_diario(cp(tipo_cp="07"), config) == "20"
    # Boletas y honorarios NO se mueven: tienen registro propio
    assert concar.sub_diario(cp(tipo_cp="03"), config) == "13"
    assert concar.sub_diario(cp(tipo_cp="02"), config) == "15"
    # La detracción sigue ganando al general
    det = cp(detraccion={"codigo": "012", "porcentaje": 10})
    assert concar.sub_diario(det, config) == "10"
    # Y un tipos.NN.sub_diario puesto por el estudio gana al general (override fino por jsonb)
    fino = configuracion({"sub_diario_compras": "20", "tipos": {"01": {"sub_diario": "77"}}})
    assert concar.sub_diario(cp(), fino) == "77" and concar.sub_diario(cp(tipo_cp="12"), fino) == "20"


def test_etiquetas_sub_diario_siguen_a_la_renumeracion():
    # Con los defaults, el mapa clásico
    assert concar.etiquetas_sub_diario(CONTAB) == {
        "05": "Ventas", "11": "Compras", "10": "Compras con detracción",
        "13": "Boletas de venta", "15": "Recibos por honorarios",
    }
    # El estudio renumera honorarios a 33: el 33 sale etiquetado, el 15 desaparece
    config = configuracion({"tipos": {"02": {"sub_diario": "33"}}})
    et = concar.etiquetas_sub_diario(config)
    assert et["33"] == "Recibos por honorarios" and "15" not in et
    # Un CONCAR que no separa la detracción: dos usos comparten número y se leen juntos
    junto = configuracion({"sub_diario_detraccion": "11"})
    assert concar.etiquetas_sub_diario(junto)["11"] == "Compras · Compras con detracción"


# ── Lo que el driver de CONCAR se niega a escribir (0.8.0) ─────────────────────

def test_el_driver_se_niega_sin_centro_donde_la_cuenta_lo_lleva():
    """La regla del contador (06-sep-2026), que hasta la 0.7 solo aplicaba el portal: obligatorio donde
    la cuenta lo lleva en la M. Decisión de John (11-sep-2026): la hace cumplir el driver."""
    with pytest.raises(concar.SinCentro) as e:
        construir_concar(COMPRAS, [cp(centro_costo="")], CONTAB, {"11": 1})
    assert [c.numero for c in e.value.comprobantes] == ["00000123"]
    # Una cuenta que no lleva centro (603201 no empieza por 63, 65 ni 70) sale igual sin él.
    contenido, _ = construir_concar(COMPRAS, [cp(cuenta_contable="603201", centro_costo="")], CONTAB, {"11": 1})
    assert contenido[:2] == b"PK"


def test_el_driver_se_niega_si_el_correlativo_pasa_de_9999():
    """CONCAR numera con MM + cuatro dígitos; el portal lo comprobaba por su cuenta hasta la 0.8."""
    with pytest.raises(driver_concar.CorrelativoDesborda) as e:
        construir_concar(COMPRAS, [cp(), cp(numero="124")], CONTAB, {"11": 9999})
    assert e.value.sub_diarios == {"11": 10000} and "supera los 4 dígitos" in str(e.value)
    contenido, resumen = construir_concar(COMPRAS, [cp()], CONTAB, {"11": 9999})
    assert contenido[:2] == b"PK" and resumen["sub_diarios"]["11"]["hasta"] == 9999


# ── CONCAR en la forma `desde_lineas` (1.0) ─────────────────────────────────────

def _celdas(contenido: bytes) -> list[tuple]:
    return list(openpyxl.load_workbook(io.BytesIO(contenido))["CONCAR"].iter_rows(values_only=True))


def test_la_tasa_de_la_ao_sale_del_comprobante_y_no_de_la_linea_redondeada():
    """IGV 175.00 sobre base 1000.03: el 17.4995 % da 17 redondeando una sola vez. Desde la tasa de la línea, que ya va
    redondeada a 17.50, daría 18. Por eso la AO se calcula con la cabecera del comprobante que trae el índice."""
    c = cp(base_gravada="1000.03", igv="175.00", total="1175.03")
    linea = concar.lineas_del_comprobante(c, CONTAB, MES, "080001")[0]
    assert Decimal(str(linea.tasa_igv)) == Decimal("17.50")
    contenido, _ = construir_concar(COMPRAS, [c], CONTAB, {"11": 1})
    assert _celdas(contenido)[3][40] == 17
    assert driver_concar.filas_de_comprobante(c, CONTAB, MES, "080001")[0]["AO"] == 17


def test_el_indice_parte_las_lineas_por_comprobante_con_su_cabecera():
    cs = [cp(), cp(numero="124"), cp(tipo_cp="03", serie="B001", numero="55", igv="0", base_gravada="0", inafecto="118")]
    lineas, rangos, indice = concar.lineas_e_indice_del_libro(COMPRAS, cs, CONTAB, {"11": 20, "13": 1})
    assert [(e.desde, e.hasta) for e in indice] == [(0, 3), (3, 6), (6, 8)] and len(lineas) == 8
    assert [(e.sub_diario, e.correlativo) for e in indice] == [("11", "080020"), ("11", "080021"), ("13", "080001")]
    assert all({ln.correlativo for ln in e.lineas(lineas)} == {e.correlativo} for e in indice)
    assert indice[0].cabecera.glosa == concar.glosa_de(cs[0]) and Decimal(indice[2].cabecera.igv) == 0
    assert indice[1].a_dict()["lineas"] == [3, 6]


def test_el_resumen_de_concar_es_el_de_la_numeracion():
    cs = [cp(), cp(numero="124"), cp(tipo_cp="03", serie="B001", numero="55", igv="0", base_gravada="0", inafecto="118")]
    _, resumen = construir_concar(COMPRAS, cs, CONTAB, {"11": 20, "13": 1})
    _, rangos = concar.numerar_en_orden(cs, CONTAB, "202608", {"11": 20, "13": 1})
    etiquetas = concar.etiquetas_sub_diario(CONTAB)
    assert resumen["sub_diarios"] == {s: {"etiqueta": etiquetas.get(s, s), **r} for s, r in rangos.items()}
    assert resumen["fechas"] == "por comprobante (extemporáneos al 01/08/2026)" and resumen["filas"] == 8


def test_sin_indice_concar_no_escribe():
    lineas = concar.lineas_del_comprobante(cp(), CONTAB, MES, "080001")
    with pytest.raises(ValueError, match="indice"):
        driver_concar.desde_lineas(COMPRAS, lineas, CONTAB)
